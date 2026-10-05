import datetime as dt
import sys
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.styles import PatternFill  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "enews_robotas"))
import nustatymai  # noqa: E402
from excel_eiles import Sarasas, menuo_diena  # noqa: E402

SIANDIEN = dt.date(2026, 10, 1)


@pytest.mark.parametrize("reiksme, laukiama", [
    ("09.24", dt.date(2026, 9, 24)),
    (9.24, dt.date(2026, 9, 24)),
    (9.1, dt.date(2026, 9, 10)),          # Excel skaičius 09.10 → 9.1
    (dt.datetime(2026, 9, 22), dt.date(2026, 9, 22)),
    ("12.30", dt.date(2025, 12, 30)),     # ateities data → praėję metai
    ("", None),
    ("abc", None),
    ("13.40", None),
])
def test_menuo_diena(reiksme, laukiama):
    assert menuo_diena(reiksme, SIANDIEN) == laukiama


def _sarasas(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    eil = [
        ("TESTA000000000001", "TST01", "1T1T01", "T0001", "09.24", "09.24", "AAA001"),
        ("TESTA000000000002", "TST02", "1T2T02", "T0002", "09.22", "09.24", "AAA002"),
        ("TESTA000000000006", "TST01", "", "T0006", "09.24", "09.25", "AA062"),
    ]
    for i, (vin, a, b, c, pdi, te, nr) in enumerate(eil, start=1):
        ws[f"B{i}"], ws[f"E{i}"], ws[f"G{i}"], ws[f"I{i}"] = vin, a, b, c
        ws[f"D{i}"] = ws[f"F{i}"] = ws[f"H{i}"] = "-"
        ws[f"K{i}"], ws[f"L{i}"], ws[f"M{i}"] = pdi, te, nr
    ws["B1"].fill = PatternFill(fill_type="solid", fgColor="FF00B050")  # jau padaryta
    ws["B4"] = "(PDI date)"
    kelias = tmp_path / "masinos.xlsx"
    wb.save(kelias)
    return kelias


def test_neapdorotos_praleidzia_nuspalvintas(tmp_path):
    s = Sarasas(_sarasas(tmp_path), nustatymai)
    masinos = s.neapdorotos(SIANDIEN)
    assert [m.eilute for m in masinos] == [2, 3]
    m = masinos[0]
    assert m.kodas_tekstu == "TST02-1T2T02-T0002"
    assert (m.pdi_data, m.garantija, m.numeris) == (dt.date(2026, 9, 22), dt.date(2026, 9, 24), "AAA002")
    assert not m.klaidos
    assert masinos[1].klaidos == ["nepilnas Midtronics kodas"]


def test_pazymeti_issaugo_spalva_ir_busena(tmp_path):
    kelias = _sarasas(tmp_path)
    s = Sarasas(kelias, nustatymai)
    s.pazymeti(2, nustatymai.SPALVA_ATLIKTA, "Atlikta")
    s.issaugoti()
    s2 = Sarasas(kelias, nustatymai)
    assert s2.ws["N2"].value == "Atlikta"
    assert [m.eilute for m in s2.neapdorotos(SIANDIEN)] == [3]


def test_tikrinimas_randa_dublikatus_ir_formatus(tmp_path):
    import tikrinimas
    s = Sarasas(_sarasas(tmp_path), nustatymai)
    s.irasyti(3, "G", "1T2T02")
    s.irasyti(3, "E", "TST02")
    s.irasyti(3, "I", "T0002")    # tas pats kodas kaip 2 eil.
    s.irasyti(2, "M", "AAA 00")   # ne ABC123
    s.irasyti(2, "K", "09.30")    # PDI vėliau nei tech. pradžia – normalu, nežymima
    pastabos = tikrinimas.duomenys(s.visos(SIANDIEN), SIANDIEN)
    tekstai = {(p.eilute, p.lygis, p.tekstas.split(" ")[0]) for p in pastabos}
    assert (2, tikrinimas.KLAIDA, "Midtronics") in tekstai
    assert (3, tikrinimas.KLAIDA, "Midtronics") in tekstai
    assert (2, tikrinimas.PERSPEJIMAS, "numeris") in tekstai
    assert (2, tikrinimas.PERSPEJIMAS, "PDI") not in tekstai
    assert tikrinimas.blokuojamos_eilutes(pastabos) == {2, 3}


def test_kodas_jau_panaudotas_padarytoje_eiluteje(tmp_path):
    import tikrinimas
    s = Sarasas(_sarasas(tmp_path), nustatymai)
    for st, v in (("E", "TST01"), ("G", "1T1T01"), ("I", "T0001")):  # kaip nuspalvintoje 1 eil.
        s.irasyti(2, st, v)
    pastabos = tikrinimas.duomenys(s.visos(SIANDIEN), SIANDIEN)
    assert any(p.eilute == 2 and "kartojasi" in p.tekstas for p in pastabos)
    assert not any(p.eilute == 1 for p in pastabos)  # padarytų eilučių nekritikuojame


def test_bloga_vin_rodoma_su_klaida(tmp_path):
    s = Sarasas(_sarasas(tmp_path), nustatymai)
    s.irasyti(2, "B", "TESTA0000000000O2")  # O vietoj 0
    m = next(m for m in s.visos(SIANDIEN) if m.eilute == 2)
    assert any("VIN" in k for k in m.klaidos)


def test_nustatymai_issaugomi_json(tmp_path, monkeypatch):
    monkeypatch.setattr(nustatymai, "JSON_FAILAS", tmp_path / "n.json")
    try:
        nauji = nustatymai.dabartines()
        nauji["RIDA"] = "7"
        nauji["TEKSTAI"] = dict(nauji["TEKSTAI"], validate="Validate | Patvirtinti")
        nustatymai.issaugoti(nauji)
        assert nustatymai.RIDA == "7"
        assert nustatymai.tekstai("validate") == ["Validate", "Patvirtinti"]
        import json
        assert json.loads((tmp_path / "n.json").read_text(encoding="utf-8")) == {
            "RIDA": "7", "TEKSTAI": {"validate": "Validate | Patvirtinti"}}
    finally:
        (tmp_path / "n.json").unlink()
        nustatymai.perkrauti()
    assert nustatymai.RIDA == "5"


def test_ascii_dalis():
    pytest.importorskip("playwright")
    import robotas
    assert robotas.ascii_dalis("Išsaugoti ir uždaryti") == "saugoti ir u"
    assert robotas.ascii_dalis("Techninės priežiūros planas") == "ros planas"
    assert robotas.ascii_dalis("Validate") == "Validate"


def test_sablono_pavyzdys_perskaitomas(tmp_path, monkeypatch):
    from excel_eiles import sukurti_sablona
    kelias = sukurti_sablona(tmp_path / "s.xlsx", nustatymai)
    assert Sarasas(kelias, nustatymai).visos() == []          # pirmas lapas tuščias
    monkeypatch.setattr(nustatymai, "EXCEL_LAPAS", "Pavyzdys")
    m = Sarasas(kelias, nustatymai).visos(SIANDIEN)
    assert [(x.vin, x.kodas_tekstu, x.numeris, x.pdi_data) for x in m][0] == (
        "TESTA000000000002", "TST02-1T2T02-T0002", "AAA002", dt.date(2026, 9, 22))
    assert not any(x.klaidos for x in m)


def test_spausdinimo_planas_ir_isdestymas():
    import spausdinimas
    assert spausdinimas.planas() == [(1, "paprastas"), (1, "lipnus"), (2, "lipnus")]
    # A4 600 dpi, 100 taškų neprintinamas kraštas: puslapis sumažinamas iki spausdinamo ploto, centruotas
    x0, y0, x1, y1 = spausdinimas.isdestymas((4960, 7016), (4760, 6816), (100, 100), (595, 842), (600, 600))
    assert (x1 - x0) <= 4760 and (y1 - y0) <= 6816 and x0 >= 0 and y0 >= 0
    assert abs((x1 - x0) / (y1 - y0) - 595 / 842) < 0.01


def test_suvestine_ir_senu_failu_trynimas(tmp_path, monkeypatch):
    pytest.importorskip("playwright")
    import os
    import time
    import robotas
    from excel_eiles import Masina
    m = Masina(3, "TESTA000000000003", ("BAD01", "1T2T02", "T0003"), None, None, "AAA003")
    t = robotas.suvestine([(m, "atlikta", "Atlikta", 125), (m, "akumas", "Akumuliatorius: Replace battery", 30)])
    assert "Atlikta: 1, atidėta: 1" in t and "2:05" in t and "AAA003" in t
    monkeypatch.setattr(robotas, "SPAUSDINTI", tmp_path)
    senas, naujas = tmp_path / "senas.pdf", tmp_path / "naujas.pdf"
    senas.write_bytes(b"x"); naujas.write_bytes(b"x")
    os.utime(senas, (time.time() - 40 * 86400,) * 2)
    assert robotas.valyti_senus(30) == 1 and not senas.exists() and naujas.exists()


def test_narsykles_pasirinkimas(tmp_path, monkeypatch):
    import tikrinimas
    brave = tmp_path / "brave.exe"
    brave.write_bytes(b"")
    monkeypatch.setattr(tikrinimas, "NARSYKLES", {"chrome": [tmp_path / "nera.exe"], "brave": [brave],
                                                  "edge": [tmp_path / "nera2.exe"]})
    assert tikrinimas.rasti_narsykle("auto") == ("brave", brave)    # Chrome nėra – imamas Brave
    assert tikrinimas.rasti_narsykle("brave") == ("brave", brave)
    assert tikrinimas.rasti_narsykle("edge") is None
