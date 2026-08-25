import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402

from app.aikstele_docx import (  # noqa: E402
    build_document,
    build_values,
    fill_template,
    suggested_filename,
)

DATA = {
    "make": "NISSAN",
    "type": "F16",
    "variant": "A",
    "version": "A45",
    "commercial_name": "NISSAN JUKE",
    "vin": "SJNF16FA7U2000002",
    "approval_number": "e9*2007/46*6697*17",
    "approval_date": "16.10.2025",
    "national_approval_number": "",
    "colour": "GELTONA",
}


def all_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_build_document_contains_all_values():
    values = build_values(DATA, doc_number="17", doc_date="2026-08-17")
    text = all_text(build_document(values, company_line="UAB Pavyzdys"))
    assert "NISSAN JUKE" in text
    assert "F16/A/A45" in text
    assert "SJNF16FA7U2000002" in text
    assert "e9*2007/46*6697*17" in text
    assert "GELTONA" in text
    assert "APIE TRANSPORTO PRIEMONĖS TAPATUMO DUOMENIS" in text


def test_fill_template_writes_into_value_column():
    empty = build_values({key: "" for key in DATA}, doc_date="")
    template = build_document(empty)
    filled, warnings = fill_template(
        template, build_values(DATA, doc_number="17", doc_date="2026-08-17")
    )
    assert warnings == []

    document = Document(io.BytesIO(filled))
    table = document.tables[-1]
    rows = {row.cells[0].text: [c.text for c in row.cells] for row in table.rows}
    assert rows["Gamybinė markė (gamintojo prekės pavadinimas):"][2] == "NISSAN"
    assert rows["Tipas/Variantas/Versija:"][2] == "F16/A/A45"
    assert rows["Komercinis pavadinimas:"][2] == "NISSAN JUKE"
    assert rows["Transporto priemonės identifikavimo numeris"][2] == "SJNF16FA7U2000002"
    assert rows["Tipo patvirtinimo Nr."][2] == "e9*2007/46*6697*17"
    assert rows["Tipo patvirtinimo numerio suteikimo data"][2] == "16.10.2025"
    assert rows["Transporto priemonės spalva"][2] == "GELTONA"
    # RL skiltis lieka nepaliesta
    assert rows["Gamybinė markė (gamintojo prekės pavadinimas):"][3] == "D.1"


def test_fill_template_fills_date_boxes():
    template = build_document(build_values({key: "" for key in DATA}, doc_date="2000-01-01"))
    filled, _ = fill_template(template, build_values(DATA, doc_date="2026-08-17"))
    document = Document(io.BytesIO(filled))
    boxes = "".join(cell.text for cell in document.tables[0].rows[0].cells)
    assert boxes == "2026-08-17"


def test_fill_template_supports_placeholders():
    document = Document()
    document.add_paragraph("Markė: {{make}}, VIN: {{vin}}, spalva: {{colour}}")
    buffer = io.BytesIO()
    document.save(buffer)

    filled, _ = fill_template(buffer.getvalue(), build_values(DATA))
    text = all_text(filled)
    assert "Markė: NISSAN" in text
    assert "VIN: SJNF16FA7U2000002" in text
    assert "spalva: GELTONA" in text


def test_fill_template_three_column_table():
    document = Document()
    table = document.add_table(rows=3, cols=3)
    for row, label, rl in zip(
        table.rows,
        ("Gamybinė markė:", "Komercinis pavadinimas:", "Transporto priemonės spalva"),
        ("D.1", "D.3", "R"),
    ):
        row.cells[0].text = label
        row.cells[2].text = rl
    buffer = io.BytesIO()
    document.save(buffer)

    filled, _ = fill_template(buffer.getvalue(), build_values(DATA))
    result = Document(io.BytesIO(filled))
    values = [row.cells[1].text for row in result.tables[0].rows]
    assert values == ["NISSAN", "NISSAN JUKE", "GELTONA"]


def test_suggested_filename_uses_vin():
    assert suggested_filename(DATA) == "aikstele_SJNF16FA7U2000002.docx"
