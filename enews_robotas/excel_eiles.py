"""Mašinų sąrašo skaitymas iš Excel ir būsenos įrašymas atgal."""

from __future__ import annotations

import datetime as dt
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string

VIN_RE = re.compile(r"^[A-HJ-NPR-Z0-9]{17}$")
# Eilutė laikoma mašina, jei B stulpelyje kažkas panašaus į VIN (net su klaida –
# tada ji parodoma su klaida, o ne tyliai praleidžiama).
PANASU_I_VIN = re.compile(r"^(?=.*\d)(?=.*[A-Z])[A-Z0-9]{14,20}$")


@dataclass
class Masina:
    eilute: int
    vin: str
    kodas: tuple[str, str, str]
    pdi_data: dt.date | None
    garantija: dt.date | None
    numeris: str
    klaidos: list[str] = field(default_factory=list)
    nuspalvinta: bool = False   # jau padaryta / atidėta – robotas nelies
    busena: str = ""            # kas parašyta būsenos stulpelyje

    @property
    def kodas_tekstu(self) -> str:
        return "-".join(self.kodas)


def menuo_diena(reiksme, siandien: dt.date) -> dt.date | None:
    """„09.24“ (tekstas, skaičius 9.24 arba Excel data) → data su metais.

    Metai – einamieji; jei data išeitų daugiau nei savaitė į ateitį
    (pvz. sausį vedama gruodžio mašina), imami praėję metai.
    """
    if reiksme is None or reiksme == "":
        return None
    if isinstance(reiksme, dt.datetime):
        reiksme = reiksme.date()
    if isinstance(reiksme, dt.date):
        menuo, diena = reiksme.month, reiksme.day
    else:
        if isinstance(reiksme, (int, float)):
            tekstas = f"{reiksme:.2f}"
        else:
            tekstas = str(reiksme).strip()
        m = re.fullmatch(r"(\d{1,2})[.,/-](\d{1,2})", tekstas)
        if not m:
            return None
        menuo, diena = int(m.group(1)), int(m.group(2))
    for metai in (siandien.year, siandien.year - 1):
        try:
            data = dt.date(metai, menuo, diena)
        except ValueError:
            return None
        if data <= siandien + dt.timedelta(days=7):
            return data
    return None


def _tekstas(reiksme) -> str:
    if reiksme is None:
        return ""
    if isinstance(reiksme, float) and reiksme.is_integer():
        reiksme = int(reiksme)
    return str(reiksme).strip().upper()


def nuspalvinta(langelis) -> bool:
    """Ar eilutė jau pažymėta spalva (žalia, oranžinė…) – tada jos neliečiame."""
    fill = langelis.fill
    if fill is None or fill.fill_type != "solid":
        return False
    spalva = fill.fgColor
    if spalva is None:
        return False
    if spalva.type == "rgb":
        return str(spalva.rgb).upper() not in ("00000000", "FFFFFFFF")
    if spalva.type == "indexed":
        return spalva.indexed not in (64, 65, 9)  # sistemos fonas / balta
    if spalva.type == "theme":
        # 0 ir 1 su tint 0 – balta / juoda fono tema
        return not (spalva.theme in (0, 1) and not spalva.tint)
    return True


class Sarasas:
    def __init__(self, kelias: str | Path, nust):
        self.kelias = Path(kelias)
        self.n = nust
        self.wb = load_workbook(self.kelias)
        self.ws = self.wb[nust.EXCEL_LAPAS] if nust.EXCEL_LAPAS else self.wb.worksheets[0]

    def _r(self, stulpelis: str, eilute: int):
        return self.ws.cell(row=eilute, column=column_index_from_string(stulpelis))

    def neapdorotos(self, siandien: dt.date | None = None) -> list[Masina]:
        return [m for m in self.visos(siandien) if not m.nuspalvinta]

    def visos(self, siandien: dt.date | None = None) -> list[Masina]:
        """Visos eilutės su VIN (ir jau nuspalvintos – jas rodo valdymo langas)."""
        siandien = siandien or dt.date.today()
        n = self.n
        masinos = []
        for eil in range(1, self.ws.max_row + 1):
            vin_l = self._r(n.STULP_VIN, eil)
            vin = _tekstas(vin_l.value)
            if not PANASU_I_VIN.match(vin):
                continue
            kodas = tuple(_tekstas(self._r(s, eil).value) for s in n.STULP_KODAS)
            m = Masina(
                eilute=eil,
                vin=vin,
                kodas=kodas,
                pdi_data=menuo_diena(self._r(n.STULP_PDI_DATA, eil).value, siandien),
                garantija=menuo_diena(self._r(n.STULP_GARANTIJA, eil).value, siandien),
                numeris=_tekstas(self._r(n.STULP_NUMERIS, eil).value).replace(" ", ""),
                nuspalvinta=nuspalvinta(vin_l),
                busena=str(self._r(n.STULP_BUSENA, eil).value or ""),
            )
            if not VIN_RE.match(vin):
                m.klaidos.append(f"VIN su klaida ({len(vin)} simb.; raidžių I, O, Q VIN nebūna)")
            if not all(kodas):
                m.klaidos.append("nepilnas Midtronics kodas")
            if m.pdi_data is None:
                m.klaidos.append("neaiški PDI data")
            if m.garantija is None:
                m.klaidos.append("neaiški tech. pradžios data")
            if not m.numeris:
                m.klaidos.append("nėra valst. numerio")
            masinos.append(m)
        return masinos

    def pazymeti(self, eilute: int, spalva: str, busena: str) -> None:
        fill = PatternFill(fill_type="solid", fgColor="FF" + spalva, bgColor="FF" + spalva)
        paskutinis = max(column_index_from_string(self.n.STULP_NUMERIS),
                         *(column_index_from_string(s) for s in self.n.STULP_KODAS))
        for st in range(1, paskutinis + 1):
            self.ws.cell(row=eilute, column=st).fill = fill
        self._r(self.n.STULP_BUSENA, eilute).value = busena

    def irasyti(self, eilute: int, stulpelis: str, reiksme) -> None:
        """Pakeičia vieną langelį (taisymas valdymo lange)."""
        self._r(stulpelis, eilute).value = reiksme

    def nuimti_spalva(self, eilute: int) -> None:
        """Kad robotas eilutę padarytų dar kartą (pvz. pakeitus akumuliatorių)."""
        paskutinis = max(column_index_from_string(self.n.STULP_NUMERIS),
                         *(column_index_from_string(s) for s in self.n.STULP_KODAS))
        for st in range(1, paskutinis + 1):
            self.ws.cell(row=eilute, column=st).fill = PatternFill(fill_type=None)
        self._r(self.n.STULP_BUSENA, eilute).value = None

    def issaugoti(self) -> None:
        self.wb.save(self.kelias)

    def atsargine_kopija(self) -> Path:
        laikas = dt.datetime.now().strftime("%Y%m%d-%H%M")
        kopija = self.kelias.with_name(f"{self.kelias.stem}.atsargine-{laikas}{self.kelias.suffix}")
        shutil.copy2(self.kelias, kopija)
        return kopija


def uzrakintas(kelias: str | Path) -> bool:
    """Ar failas atidarytas Excel'yje (Windows tada neleidžia į jį rašyti)."""
    try:
        with open(kelias, "r+b"):
            return False
    except PermissionError:
        return True


def sukurti_sablona(kelias: str | Path, nust) -> Path:
    """Tuščias Excel šablonas su antraštėmis tose vietose, kur robotas skaito.
    Antrame lape „Pavyzdys“ – kaip užpildyti (robotas skaito tik pirmą lapą)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    k1, k2, k3 = nust.STULP_KODAS
    antrastes = {
        nust.STULP_VIN: "VIN", k1: "Kodas 1", k2: "Kodas 2", k3: "Kodas 3",
        nust.STULP_PDI_DATA: "PDI data (MM.DD)", nust.STULP_GARANTIJA: "Tech. pradžia (MM.DD)",
        nust.STULP_NUMERIS: "Valst. nr.", nust.STULP_BUSENA: "Būsena (pildo robotas)",
    }
    pavyzdziai = [
        ("TESTA000000000002", "TST02", "1T2T02", "T0002", "09.22", "09.24", "AAA002"),
        ("TESTA000000000003", "TST01", "1T1T01", "T0001", "09.24", "09.25", "AAA003"),
    ]
    tekstiniai = [nust.STULP_VIN, k1, k2, k3, nust.STULP_PDI_DATA, nust.STULP_GARANTIJA, nust.STULP_NUMERIS]

    wb = Workbook()
    for lapas, eilutes in ((wb.active, []), (wb.create_sheet("Pavyzdys"), pavyzdziai)):
        if not eilutes:
            lapas.title = "Mašinos"
        for st, tekstas in antrastes.items():
            c = lapas[f"{st}1"]
            c.value = tekstas
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, horizontal="center")
        for i in range(2, 300):  # tekstinis formatas, kad Excel nepaverstų 09.10 į 9.1
            for st in tekstiniai:
                lapas[f"{st}{i}"].number_format = "@"
        for i, (vin, a, b, c, pdi, te, nr) in enumerate(eilutes, start=2):
            for st, v in ((nust.STULP_VIN, vin), (k1, a), (k2, b), (k3, c),
                          (nust.STULP_PDI_DATA, pdi), (nust.STULP_GARANTIJA, te), (nust.STULP_NUMERIS, nr)):
                lapas[f"{st}{i}"] = v
        for n in range(1, 15):
            lapas.column_dimensions[get_column_letter(n)].width = 4
        lapas.column_dimensions[nust.STULP_VIN].width = 22
        for st in (k1, k2, k3, nust.STULP_NUMERIS):
            lapas.column_dimensions[st].width = 10
        for st in (nust.STULP_PDI_DATA, nust.STULP_GARANTIJA):
            lapas.column_dimensions[st].width = 13
        lapas.column_dimensions[nust.STULP_BUSENA].width = 40
        lapas.row_dimensions[1].height = 30
        lapas.freeze_panes = "A2"
    kelias = Path(kelias)
    kelias.parent.mkdir(parents=True, exist_ok=True)
    wb.save(kelias)
    return kelias
