"""Bandomasis eNEWS: robotą galima išbandyti namuose, be Nissan B2B.

Kompiuteryje paleidžiamas mažas serveris su eNEWS kopija (pagal tikro eNEWS
nuotraukas). Ji elgiasi kaip tikras eNEWS:
  * kodas, prasidedantis „BAD“ (pvz. BAD36-1Q9D77-TE204) → „Replace battery“;
  * kodas, jau panaudotas kitai mašinai → „test code already used“;
  * nepažymėjus visų PDI punktų ar neįvedus datos – PDI neišsaugomas;
  * Drukāt duoda PDF su mašinos duomenimis.
"""

from __future__ import annotations

import datetime as dt
import shutil
import threading
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

ARCH = Path(__file__).resolve().parent / "bandymas"


def pdf(eilutes: list[str]) -> bytes:
    """Paprasčiausias vieno puslapio PDF su tekstu (be papildomų bibliotekų)."""
    tekstas = "BT /F1 14 Tf 60 780 Td 20 TL " + " ".join(
        "(" + e.encode("latin-1", "replace").decode("latin-1").replace("\\", "\\\\")
        .replace("(", "\\(").replace(")", "\\)") + ") '" for e in eilutes) + " ET"
    objektai = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(tekstas)} >>\nstream\n{tekstas}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    rez = b"%PDF-1.4\n"
    poslinkiai = []
    for i, o in enumerate(objektai, start=1):
        poslinkiai.append(len(rez))
        rez += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(rez)
    rez += f"xref\n0 {len(objektai) + 1}\n0000000000 65535 f \n".encode()
    rez += "".join(f"{p:010d} 00000 n \n" for p in poslinkiai).encode()
    rez += f"trailer\n<< /Size {len(objektai) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return rez


class Tvarkytojas(SimpleHTTPRequestHandler):
    def do_GET(self):
        u = urlparse(self.path)
        if u.path == "/dokumentas.pdf":
            q = {k: v[0] for k, v in parse_qs(u.query).items()}
            pavadinimas = ("WARRANTY CERTIFICATE" if q.get("tipas") == "sert"
                           else "TECHNINES PRIEZIUROS PLANAS")
            turinys = pdf([f"BANDYMAS - {pavadinimas}", "",
                           f"VIN: {q.get('vin', '')}", f"Reg. No: {q.get('reg', '')}",
                           f"Warranty Start Date: {q.get('wsd', '')}",
                           f"Sukurta: {dt.datetime.now():%Y-%m-%d %H:%M}"])
            self.send_response(200)
            self.send_header("Content-Type", "application/pdf")
            if q.get("tipas") != "sert":  # antras failas – kaip atsisiuntimas (abu būdai išbandomi)
                self.send_header("Content-Disposition", f'attachment; filename="TP_{q.get("vin", "")}.pdf"')
            self.send_header("Content-Length", str(len(turinys)))
            self.end_headers()
            self.wfile.write(turinys)
            return
        if u.path in ("/", ""):
            self.send_response(302)
            self.send_header("Location", "/enews/hp_new.html")
            self.end_headers()
            return
        super().do_GET()

    def log_message(self, *args):  # tyliai
        pass


def paleisti() -> tuple[ThreadingHTTPServer, str]:
    """Paleidžia serverį laisvame prievade. Grąžina (serveris, pradžios adresas).
    Kiekvieną kartą – naujas prievadas, todėl ir švari eNEWS „atmintis“."""
    serveris = ThreadingHTTPServer(("127.0.0.1", 0), partial(Tvarkytojas, directory=str(ARCH)))
    threading.Thread(target=serveris.serve_forever, daemon=True).start()
    return serveris, f"http://127.0.0.1:{serveris.server_address[1]}/enews/hp_new.html"


def paruosti_excel(tikras: str | Path) -> Path:
    """Bandymui – tikro Excel kopija (tikras failas nepaliečiamas), o jei jo nėra – pavyzdys."""
    kopija = ARCH / "bandymui.xlsx"
    if Path(tikras).is_file():
        shutil.copy2(tikras, kopija)
    else:
        sukurti_pavyzdi(kopija)
    return kopija


def sukurti_pavyzdi(kelias: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    wb = Workbook()
    ws = wb.active
    eilutes = [
        # VIN, kodas (3 dalys), PDI, tech. pradžia, numeris, jau padaryta?
        ("SJNJ12TD3U2389917", "JRJ36", "1Q1H77", "S3604", "09.24", "09.24", "OAU413", True),
        ("SJNJ12TD0U2373741", "JRH36", "1Q9D77", "TE204", "09.22", "09.24", "OAU289", False),
        ("SJNJ12TD5U2387571", "BAD36", "1Q9D77", "SA004", "09.22", "09.22", "AYT599", False),  # blogas akumas
        ("SJNJ12TD2U2378584", "JRK36", "1Q1578", "SA804", "09.22", "09.22", "AYT592", False),
        ("SJNJ12TD6U2377406", "JRJ36", "", "T4204", "09.22", "09.23", "OAU052", False),      # trūksta kodo dalies
        ("SJNJ12TD3U2404660", "JRJ36", "1Q1H77", "TE504", "09.28", "09.25", "OAU57", False),  # keistas numeris
    ]
    for i, (vin, a, b, c, pdi, te, nr, padaryta) in enumerate(eilutes, start=1):
        ws[f"B{i}"], ws[f"D{i}"], ws[f"E{i}"], ws[f"F{i}"], ws[f"G{i}"] = vin, "-", a, "-", b
        ws[f"H{i}"], ws[f"I{i}"], ws[f"K{i}"], ws[f"L{i}"], ws[f"M{i}"] = "-", c, pdi, te, nr
        if padaryta:
            for st in "ABCDEFGHIJKLM":
                ws[f"{st}{i}"].fill = PatternFill(fill_type="solid", fgColor="FF00B050")
            ws[f"N{i}"] = "Atlikta (pavyzdys)"
    ws[f"K{len(eilutes) + 1}"], ws[f"L{len(eilutes) + 1}"] = "(PDI date)", "(Te. Date)"
    kelias.parent.mkdir(parents=True, exist_ok=True)
    wb.save(kelias)
    return kelias


if __name__ == "__main__":
    import webbrowser
    _, adresas = paleisti()
    print(f"Bandomasis eNEWS: {adresas}  (Ctrl+C – baigti)")
    webbrowser.open(adresas)
    threading.Event().wait()
