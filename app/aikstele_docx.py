"""Pažymos apie transporto priemonės tapatumo duomenis ("aikštelės") kūrimas.

Du režimai:

1. **Šablono režimas** (rekomenduojamas) – imamas Jūsų realus Word failas ir
   užpildomos jo lentelės eilutės. Šablone nieko keisti nereikia: eilutės
   randamos pagal pirmame stulpelyje esantį pavadinimą ("Gamybinė markė...",
   "Tipas/Variantas/Versija" ir t. t.). Papildomai palaikomi ir `{{žymekliai}}`.

2. **Kūrimo iš naujo režimas** – dokumentas suformuojamas nuo nulio pagal
   pažymos struktūrą (naudokite, jei šablono po ranka nėra).
"""

from __future__ import annotations

import io
import unicodedata
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

# ---------------------------------------------------------------------------
# Laukų aprašai
# ---------------------------------------------------------------------------

#: (raktas, pažymos pavadinimas, CoC skirsnis, RL skiltis)
PAZYMA_ROWS: list[tuple[str, str, str, str]] = [
    ("make", "Gamybinė markė (gamintojo prekės pavadinimas):", "0.1", "D.1"),
    ("type_variant_version", "Tipas/Variantas/Versija:", "0.2", "D.2"),
    ("commercial_name", "Komercinis pavadinimas:", "0.2.1", "D.3"),
    ("vin", "Transporto priemonės identifikavimo numeris", "0.10", "E"),
    ("approval_number", "Tipo patvirtinimo Nr.", "", "K"),
    ("approval_date", "Tipo patvirtinimo numerio suteikimo data", "", ""),
    ("national_approval_number", "Nacionalinis patvirtinimo numeris", "", "K1"),
    ("colour", "Transporto priemonės spalva", "40", "R"),
]

#: Pagal ką atpažįstamos šablono lentelės eilutės (tikrinama iš eilės).
ROW_MATCHERS: list[tuple[str, tuple[str, ...]]] = [
    ("national_approval_number", ("nacionalinis patvirtinimo numeris", "nacionalinis")),
    ("approval_date", ("suteikimo data", "patvirtinimo data")),
    ("approval_number", ("tipo patvirtinimo nr", "tipo patvirtinimo numeris")),
    ("vin", ("identifikavimo numeris", "vin")),
    ("commercial_name", ("komercinis pavadinimas",)),
    ("type_variant_version", ("tipas/variantas/versija", "tipas / variantas", "tipas")),
    ("make", ("gamybine marke", "marke")),
    ("colour", ("spalva",)),
]

PASTABA = (
    "Pastaba. Duomenys pildomi tik tuo atveju, kai jie yra nurodyti atitikties "
    "liudijime."
)

# PASTABA: šis tekstas perrašytas iš pažymos pavyzdžio ir gali skirtis nuo
# Jūsų įmonėje naudojamos redakcijos. Naudojant šabloną (1 režimas) jis
# imamas iš Jūsų dokumento ir čia esantis tekstas nenaudojamas.
PATVIRTINIMAS = (
    "Patvirtinu, kad transporto priemonės identifikavimo žymenys atitinka "
    "Komisijos reglamento (ES) Nr. 19/2011 dėl variklinių transporto priemonių "
    "ir jų priekabų tipo patvirtinimo, atsižvelgiant į gamintojo identifikavimo "
    "plokštelę ir transporto priemonės identifikavimo numerį, reikalavimus, "
    "kuriuo įgyvendinamas Europos Parlamento ir Tarybos reglamentas (EB) "
    "Nr. 661/2009 dėl variklinių transporto priemonių, jų priekabų ir joms "
    "skirtų sistemų, sudėtinių dalių bei atskirų techninių mazgų."
)


# ---------------------------------------------------------------------------
# Pagalbinės funkcijos
# ---------------------------------------------------------------------------


def _fold(text: str) -> str:
    """Normalizuoja tekstą palyginimui: be diakritikų, mažosiomis, be kabučių."""
    text = unicodedata.normalize("NFKD", text or "")
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().replace("\xa0", " ")
    text = "".join(ch if (ch.isalnum() or ch in " /") else " " for ch in text)
    return " ".join(text.split())


def _set_cell_text(cell, value: str) -> None:
    """Įrašo reikšmę į langelį išsaugant pirmo teksto fragmento formatavimą."""
    paragraphs = cell.paragraphs
    para = paragraphs[0]
    for extra in paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    if para.runs:
        para.runs[0].text = value
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(value)


def _cell_text(cell) -> str:
    return " ".join(p.text for p in cell.paragraphs).strip()


def _row_cells(row):
    """Eilutės langeliai be pasikartojimų (sujungti langeliai kartojasi)."""
    seen: list = []
    out = []
    for cell in row.cells:
        if any(cell._element is s for s in seen):
            continue
        seen.append(cell._element)
        out.append(cell)
    return out


def build_values(data: dict, doc_number: str = "", doc_date: str = "") -> dict:
    """Iš CoC duomenų žodyno paruošia visas dokumento reikšmes."""
    values = dict(data)
    if not values.get("type_variant_version"):
        parts = [values.get(k, "") for k in ("type", "variant", "version")]
        values["type_variant_version"] = "/".join(p for p in parts if p)
    values["doc_number"] = doc_number or ""
    values["doc_date"] = doc_date or date.today().isoformat()
    values["today"] = date.today().strftime("%Y-%m-%d")
    values.setdefault("national_approval_number", "")
    return values


# ---------------------------------------------------------------------------
# 1 režimas: esamo šablono pildymas
# ---------------------------------------------------------------------------


def _replace_placeholders(paragraph, values: dict) -> bool:
    """Pakeičia `{{raktas}}` žymeklius pastraipoje. Grąžina True, jei keitė."""
    full = "".join(run.text for run in paragraph.runs)
    if "{{" not in full:
        return False
    new = full
    for key, value in values.items():
        new = new.replace("{{" + key + "}}", str(value or ""))
    if new == full:
        return False
    paragraph.runs[0].text = new
    for run in paragraph.runs[1:]:
        run.text = ""
    return True


def _find_value_column(table, matched_rows: list[tuple[int, str]]) -> int | None:
    """Nustato, kuriame stulpelyje rašomos reikšmės.

    Reikšmių stulpelis yra tas, kuris atpažintose eilutėse dažniausiai tuščias
    (neskaitant paskutinio "Skiltis RL" stulpelio ir pirmojo – pavadinimų).
    """
    if not matched_rows:
        return None
    rows = [_row_cells(table.rows[i]) for i, _ in matched_rows]
    width = min(len(r) for r in rows)
    if width < 2:
        return None
    candidates = range(1, width - 1) if width > 2 else range(1, width)
    best, best_score = None, -1
    for col in candidates:
        empty = sum(1 for r in rows if not _cell_text(r[col]))
        if empty >= best_score:  # lygiosiomis renkamės dešiniausią
            best, best_score = col, empty
    return best


def _fill_tables(document, values: dict) -> list[str]:
    """Užpildo lenteles pagal eilučių pavadinimus. Grąžina užpildytus raktus."""
    filled: list[str] = []
    for table in document.tables:
        matched: list[tuple[int, str]] = []
        for idx, row in enumerate(table.rows):
            cells = _row_cells(row)
            if len(cells) < 2:
                continue
            label = _fold(_cell_text(cells[0]))
            if not label:
                continue
            for key, needles in ROW_MATCHERS:
                if key in (k for _, k in matched):
                    continue
                if any(_fold(n) in label for n in needles):
                    matched.append((idx, key))
                    break
        if len(matched) < 2:
            continue

        col = _find_value_column(table, matched)
        if col is None:
            continue
        for idx, key in matched:
            cells = _row_cells(table.rows[idx])
            if col >= len(cells):
                continue
            value = str(values.get(key, "") or "")
            if not value:
                continue
            _set_cell_text(cells[col], value)
            filled.append(key)
    return filled


def _fill_date_boxes(document, doc_date: str) -> bool:
    """Užpildo datos langelius (po vieną simbolį), jei tokia lentelė yra."""
    digits = [ch for ch in doc_date]
    for table in document.tables:
        if len(table.rows) != 1:
            continue
        cells = _row_cells(table.rows[0])
        if not (8 <= len(cells) <= 12):
            continue
        if any(len(_cell_text(c)) > 1 for c in cells):
            continue
        if len(digits) != len(cells):
            continue
        for cell, ch in zip(cells, digits):
            _set_cell_text(cell, ch)
        return True
    return False


def _fill_number(document, doc_number: str) -> bool:
    """Įrašo pažymos numerį į pastraipą, prasidedančią "Nr."."""
    if not doc_number:
        return False
    for para in document.paragraphs:
        text = para.text.strip()
        if len(text) <= 40 and text.lower().startswith("nr."):
            rest = text[3:].strip(" _. ")
            if rest:
                return False  # numeris jau įrašytas – neliečiame
            if para.runs:
                para.runs[0].text = f"Nr. {doc_number}"
                for run in para.runs[1:]:
                    run.text = ""
            else:
                para.add_run(f"Nr. {doc_number}")
            return True
    return False


def fill_template(template: str | Path | bytes, values: dict) -> tuple[bytes, list[str]]:
    """Užpildo esamą pažymos šabloną. Grąžina (docx baitai, įspėjimai)."""
    source = io.BytesIO(template) if isinstance(template, (bytes, bytearray)) else str(template)
    document = Document(source)
    warnings: list[str] = []

    replaced_any = False
    for para in document.paragraphs:
        replaced_any |= _replace_placeholders(para, values)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replaced_any |= _replace_placeholders(para, values)
    for section in document.sections:
        for container in (section.header, section.footer):
            for para in container.paragraphs:
                replaced_any |= _replace_placeholders(para, values)

    filled = _fill_tables(document, values)
    _fill_date_boxes(document, values.get("doc_date", ""))
    _fill_number(document, values.get("doc_number", ""))

    expected = {key for key, *_ in PAZYMA_ROWS if values.get(key)}
    missing = expected - set(filled)
    if missing and not replaced_any:
        names = {key: name for key, name, *_ in PAZYMA_ROWS}
        warnings.append(
            "Šablone nerastos eilutės: "
            + ", ".join(sorted(names[k].rstrip(":") for k in missing))
        )

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), warnings


# ---------------------------------------------------------------------------
# 2 režimas: dokumento kūrimas iš naujo
# ---------------------------------------------------------------------------


def _style_document(document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(0)


def _add_paragraph(document, text, *, bold=False, align=None, space_after=0, size=None):
    para = document.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if align is not None:
        para.alignment = align
    para.paragraph_format.space_after = Pt(space_after)
    return para


def build_document(values: dict, company_line: str = "") -> bytes:
    """Suformuoja pažymą nuo nulio (kai šablonas nepateiktas)."""
    document = Document()
    _style_document(document)

    _add_paragraph(document, company_line, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_paragraph(
        document,
        "(įmonės kodas, adresas)",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
        size=9,
    )

    _add_paragraph(document, "PAŽYMA", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_paragraph(
        document,
        "APIE TRANSPORTO PRIEMONĖS TAPATUMO DUOMENIS",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _add_paragraph(
        document,
        "(M, N ir O kategorijų transporto priemonei)",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )

    # datos langeliai
    doc_date = values.get("doc_date", "")
    boxes = document.add_table(rows=1, cols=max(len(doc_date), 1))
    boxes.style = "Table Grid"
    boxes.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, ch in zip(_row_cells(boxes.rows[0]), doc_date):
        cell.width = Cm(0.7)
        _set_cell_text(cell, ch)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_paragraph(document, "(data)", align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
    _add_paragraph(
        document,
        f"Nr. {values.get('doc_number', '')}".rstrip(),
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )

    # pagrindinė lentelė
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    header = _row_cells(table.rows[0])
    for cell, text in zip(
        header,
        ("Transporto priemonės duomenys:", "Skirsnis atitikties liudijime", "", "Skiltis RL"),
    ):
        _set_cell_text(cell, text)
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True

    for key, label, section, rl in PAZYMA_ROWS:
        cells = _row_cells(table.add_row())
        _set_cell_text(cells[0], label)
        _set_cell_text(cells[1], section)
        _set_cell_text(cells[2], str(values.get(key, "") or ""))
        _set_cell_text(cells[3], rl)

    _add_paragraph(document, "", space_after=12)
    _add_paragraph(document, PASTABA, bold=True, space_after=12)
    para = _add_paragraph(document, PATVIRTINIMAS, bold=True, space_after=24)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    _add_paragraph(document, "_______________________        _______________________")
    _add_paragraph(document, "(pareigų pavadinimas)                  (parašas)", size=9)
    _add_paragraph(document, "(vardas, pavardė)", size=9)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def suggested_filename(values: dict) -> str:
    """Failo pavadinimas iš VIN, pvz. `aikstele_SJNJ12TD3U2000001.docx`."""
    vin = (values.get("vin") or "be_vin").strip().replace(" ", "")
    return f"aikstele_{vin}.docx"
